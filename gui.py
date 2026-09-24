"""
China VAT Refund - Workspace de extraccion y reconciliacion.

Correr:  python src/gui.py
"""

import io
import queue
import threading
import traceback
from pathlib import Path

import FreeSimpleGUI as sg
import pymupdf
import yaml
from PIL import Image

import extract
import reconciliation
from review_store import ReviewStore

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
DEFAULT_SAMPLES = ROOT / "samples"
CORRECTIONS = ROOT / "output" / "corrections.json"

# un campo por debajo de esto se manda a revision aunque se haya encontrado
CONF_THRESHOLD = 0.85

PREVIEW_W, PREVIEW_H = 620, 300

C = {"primary": "#2563EB", "ok": "#15803D", "warn": "#D97706",
     "bad": "#DC2626", "bg": "#F4F6F8", "surface": "#FFFFFF",
     "text": "#1F2937", "muted": "#6B7280", "sel": "#DBEAFE"}

sg.theme_add_new("China", {
    "BACKGROUND": C["bg"], "TEXT": C["text"], "INPUT": C["surface"],
    "TEXT_INPUT": C["text"], "SCROLL": "#CBD5E1",
    "BUTTON": ("#FFFFFF", C["primary"]), "PROGRESS": (C["primary"], "#E5E7EB"),
    "BORDER": 1, "SLIDER_DEPTH": 0, "PROGRESS_DEPTH": 0})
sg.theme("China")


# =============================================================================
# Datos
# =============================================================================
def flatten(results):
    """
    Resultados anidados -> filas planas para la tabla.
    [archivo, pag, tipo de documento, campo, valor, conf, estado]
    """
    rows = []
    for r in results:
        if not r.get("ok"):
            rows.append([r["path"], "", "ERROR", "", r.get("error", ""), 0, "ERROR"])
            continue
        for p in r["pages"]:
            st = p["classification"]["status"]
            if st == "SKIPPED":
                # visible a proposito: si algun dia un dato aparece en una
                # pagina que salteamos, tiene que poder verse que la salteamos
                rows.append([r["path"], p["page"], "SKIPPED", "",
                             p.get("triage", {}).get("reason", ""), 0, "SKIPPED"])
                continue
            tipo = p["classification"]["type"] or "UNIDENTIFIED"
            if not p["fields"]:
                rows.append([r["path"], p["page"], tipo, "", "",
                             p["classification"]["score"], st])
                continue
            for fname, f in p["fields"].items():
                rows.append([r["path"], p["page"], tipo, fname,
                             fmt_value(f["value"]), round(f["conf"], 2), f["status"]])
    return rows


def fmt_value(v):
    if isinstance(v, dict):
        return " | ".join(f"{k}={x}" for k, x in v.items())
    if isinstance(v, list):
        return " ; ".join(map(str, v))
    return "" if v is None else str(v)


def needs_review(f):
    """Un campo va a la cola si no se encontro, quedo parcial, o la confianza
    no alcanza. Un valor con conf 0.6 no es un valor: es una sospecha."""
    if f["status"] in ("NOT_FOUND", "PARTIAL", "CONFLICT"):
        return True
    return f["status"] != "CORRECTED" and f["conf"] < CONF_THRESHOLD


def build_review_cases(results):
    cases = []
    for r in results:
        if not r.get("ok"):
            continue
        for p in r["pages"]:
            for fname, f in p["fields"].items():
                if needs_review(f):
                    cases.append({
                        "doc": r["path"], "abs_path": r.get("abs_path"),
                        "page": p["page"], "render_scale": p.get("render_scale", 1.0),
                        "rotation": p.get("rotation", 0),
                        "field": fname, "desc_zh": f.get("desc_zh", ""),
                        "value": f["value"], "conf": f["conf"],
                        "status": f["status"], "origin": f.get("origin"),
                        "doc_type": p["classification"]["type"] or "?",
                    })
    # los que no se encontraron primero: son los que bloquean
    cases.sort(key=lambda c: (c["status"] != "NOT_FOUND", c["conf"]))
    return cases


# nombres legibles para los doc_type que aparecen en 'ref' -- lo que no
# esta acá cae al fallback (snake_case -> Title Case) en _leg_label(). El
# numero entre parentesis es el prefijo del archivo en samples/ (ej. "3.2
# Exports Customs Declaration Form.pdf") -- para que el analista sepa a
# que archivo ir sin tener que adivinar por el nombre del doc_type.
DOC_LABELS = {
    "customs_declaration": "Customs Declaration (3.2)",
    "export_invoice": "Export Invoice (3.3)",
    "tax_filing_directory": "Refund Filing Directory (3.1)",
    "bill_of_lading": "Bill of Lading (3.7.1)",
    "cargo_manifest": "Cargo Manifest (3.7.2)",
    "cargo_manifest_detail": "Cargo Manifest (3.7.2)",
    "shipping_order": "Shipping Order (3.7.4.3)",
    "sales_contract": "Sales Contract (3.5)",
}


def _leg_label(leg):
    if leg["ref"]:
        doc = leg["ref"].split(".")[0]
        label = DOC_LABELS.get(doc, doc.replace("_", " ").title())
        item = leg.get("item")
        # item Code de rules.xlsx (1-84) -- para volver directo a esa fila
        # de la planilla original si hace falta reconfirmar la regla
        return f"{label}({item})" if item is not None else label
    return f"item {leg['item']}"


def _leg_value(leg):
    """Igual que fmt_value pero sin '.0' en enteros y con la unidad pegada
    -- para que 51000.0/kg se lea '51000 kg' en vez de '51000.0'."""
    v = leg["value"]
    if isinstance(v, list):
        text = ", ".join(str(x) for x in v)
    elif isinstance(v, dict):
        text = ", ".join(f"{k}={x}" for k, x in v.items())
    else:
        if isinstance(v, float) and v.is_integer():
            v = int(v)
        text = "" if v is None else str(v)
    unit = leg.get("unit")
    return f"{text} {unit}" if unit else text


def _leg_summary(leg):
    """Una pata de una regla, resumida para la columna 'Detail' de la tabla."""
    label = _leg_label(leg)
    if leg["status"] == "MISSING":
        return f"{label}: not available"
    note = {"LOSSY": "precision lost", "LOW_CONF": "low confidence"}.get(leg["status"])
    value = _leg_value(leg)
    return f"{label}: {value} — {note}" if note else f"{label}: {value}"


def flatten_reconciliation(shipments):
    """Salida de reconciliation.run() -> filas planas para -RECON_TABLE-."""
    rows = []
    for s in shipments:
        for r in s["rules"]:
            detalle = "   |   ".join(_leg_summary(leg) for leg in r["legs"])
            rows.append([s["shipment"], r["id"], r["desc"].strip(),
                        r["verdict"], detalle])
    return rows


def _truncate(text, n):
    """Recorta texto para que Description/Detail no se coman todo el ancho
    de la tabla -- el Treeview no hace wrap de multilinea, asi que la unica
    forma de acortar la celda es cortando el texto (el valor completo se ve
    abajo, en -RECON_DETAIL-, al hacer click)."""
    text = str(text)
    return text if len(text) <= n else text[:n - 1] + "…"


def _truncate_recon_row(row):
    return [row[0], row[1], _truncate(row[2], 40), row[3], _truncate(row[4], 70)]


# =============================================================================
# Preview del documento
# =============================================================================
def render_region(abs_path, page_no, origin, render_scale, rotation=0):
    """
    Devuelve el PNG de la zona del documento de donde salio el valor.
    Sin esto la revision manual es inutil: el analista tendria que abrir
    los PDFs por su cuenta para saber que esta corrigiendo.
    """
    if not abs_path or not Path(abs_path).is_file():
        return None
    try:
        doc = pymupdf.open(abs_path)
        page = doc[page_no - 1]

        if rotation:
            # origin quedo en pixeles del render YA rotado (el mismo que
            # uso extract.py para el OCR de esta pagina de costado) --
            # reproducir ese giro y recortar en espacio de pixeles evita
            # tener que reproyectar coordenadas entre marcos distintos
            # (el clip de pymupdf se interpreta ANTES de rotar, no despues).
            mat = pymupdf.Matrix(extract.OCR_DPI / 72,
                                 extract.OCR_DPI / 72).prerotate(rotation)
            img = Image.open(io.BytesIO(page.get_pixmap(matrix=mat).tobytes("png")))
            if origin:
                xs = [o["x"] for o in origin]
                ys = [o["y"] for o in origin]
                xe = [o["x"] + o["w"] for o in origin]
                ye = [o["y"] + o["h"] for o in origin]
                pad = 40
                box = (max(0, min(xs) - pad), max(0, min(ys) - pad),
                      min(img.width, max(xe) + pad), min(img.height, max(ye) + pad))
            else:
                box = (0, 0, img.width, img.height)
            crop = img.crop(tuple(int(v) for v in box))
            zoom = min(PREVIEW_W / max(crop.width, 1), PREVIEW_H / max(crop.height, 1))
            crop = crop.resize((max(1, int(crop.width * zoom)),
                               max(1, int(crop.height * zoom))))
            buf = io.BytesIO()
            crop.save(buf, format="PNG")
            doc.close()
            return buf.getvalue()

        if origin:
            # origin -> pixeles a OCR_DPI -> puntos PDF
            k = render_scale / (extract.OCR_DPI / 72)
            xs = [o["x"] * k for o in origin]
            ys = [o["y"] * k for o in origin]
            xe = [(o["x"] + o["w"]) * k for o in origin]
            ye = [(o["y"] + o["h"]) * k for o in origin]
            pad = 40
            clip = pymupdf.Rect(max(0, min(xs) - pad), max(0, min(ys) - pad),
                                min(page.rect.x1, max(xe) + pad),
                                min(page.rect.y1, max(ye) + pad))
        else:
            clip = page.rect                       # no hubo match: pagina entera

        zoom = min(PREVIEW_W / max(clip.width, 1), PREVIEW_H / max(clip.height, 1))
        pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), clip=clip)
        data = pix.tobytes("png")
        doc.close()
        return data
    except Exception:
        return None


# =============================================================================
# Layout
# =============================================================================
def tab_process():
    return [
        [sg.Frame("Documents", [[
            sg.Text("Folder", size=(10, 1)),
            sg.Input(str(DEFAULT_SAMPLES), key="-FOLDER-", expand_x=True),
            sg.FolderBrowse("Browse", target="-FOLDER-")]], expand_x=True)],
        [sg.Push(),
         sg.Button("Process", key="-PROCESS-", size=(18, 2),
                   font=("Segoe UI", 10, "bold"))],
        [sg.Text("●", key="-ICON-", text_color=C["muted"], font=("Segoe UI", 14)),
         sg.Text("Ready", key="-STATUS-", text_color=C["muted"],
                 font=("Segoe UI", 10, "bold"), size=(60, 1)),
         sg.Push(), sg.Text("0%", key="-PCT-", text_color=C["muted"])],
        [sg.ProgressBar(100, "h", size=(55, 18), key="-BAR-", expand_x=True)],
        [sg.Text("Log", font=("Segoe UI", 11, "bold"))],
        [sg.Multiline(size=(100, 13), key="-LOG-", autoscroll=True, disabled=True,
                      expand_x=True, expand_y=True, background_color="#111827",
                      text_color="#E5E7EB", font=("Consolas", 9),
                      pad=(0, (3, 20)))],
    ]


def counter(key, label, color=None):
    return sg.Frame("", [[sg.Text("0", key=key, font=("Segoe UI", 18, "bold"),
                                  text_color=color or C["text"])],
                         [sg.Text(label, text_color=C["muted"])]],
                    size=(165, 85), element_justification="center")


def tab_results():
    return [
        [counter("-N_DOCS-", "Documents"),
         counter("-N_PAGES-", "Pages"),
         counter("-N_OK-", "OK fields", C["ok"]),
         counter("-N_REV-", "To review", C["warn"]),
         counter("-N_UNK-", "Unidentified", C["bad"]), sg.Push()],
        [sg.Text("Status"),
         sg.Combo(["All", "OK", "OK_CORROBORATED", "CORRECTED", "NOT_FOUND",
                   "PARTIAL", "UNRESOLVED", "ERROR"], default_value="All",
                  readonly=True, key="-FILTER-", enable_events=True, size=(20, 1)),
         sg.Text("Search"), sg.Input(key="-SEARCH-", enable_events=True, size=(30, 1)),
         sg.Push(), sg.Button("Export to Excel", key="-EXPORT-", disabled=True,
                              button_color=("#FFFFFF", C["ok"]))],
        [sg.Table(values=[], key="-TABLE-",
                  headings=["File", "Pg", "Document type", "Field",
                            "Value", "Conf", "Status"],
                  col_widths=[26, 4, 22, 20, 34, 6, 16],
                  auto_size_columns=False, justification="left", num_rows=12,
                  alternating_row_color="#F8FAFC",
                  selected_row_colors=(C["text"], C["sel"]),
                  expand_x=True, expand_y=True, pad=(0, (3, 24)))],
    ]


def tab_review():
    left = sg.Column([
        [sg.Text("Pending", font=("Segoe UI", 12, "bold")), sg.Push(),
         sg.Text("0", key="-REV_N-", text_color="#FFFFFF",
                 background_color=C["warn"], font=("Segoe UI", 10, "bold"),
                 pad=(8, 4))],
        [sg.Listbox([], size=(40, 18), key="-REV_LIST-", enable_events=True,
                    expand_x=True, expand_y=True, pad=(0, (3, 20)))],
    ], expand_x=True, expand_y=True)

    right = sg.Column([
        [sg.Text("Select a pending case.", key="-REV_EMPTY-",
                 text_color=C["muted"], font=("Segoe UI", 11))],
        [sg.Text("Document", text_color=C["muted"]),
         sg.Text("", key="-REV_DOC-", font=("Segoe UI", 10, "bold"), size=(34, 1)),
         sg.Text("Pg", text_color=C["muted"]), sg.Text("", key="-REV_PAGE-"),
         sg.Text("Conf", text_color=C["muted"]), sg.Text("", key="-REV_CONF-")],
        [sg.Text("Field", text_color=C["muted"]),
         sg.Text("", key="-REV_FIELD-", font=("Segoe UI", 10, "bold")),
         sg.Text("", key="-REV_ZH-", text_color=C["muted"]),
         sg.Push(),
         sg.Text("", key="-REV_STATUS-", font=("Segoe UI", 10, "bold"),
                 text_color=C["bad"])],
        [sg.Frame("Document region", [[
            sg.Column([[sg.Image(key="-REV_IMG-", size=(PREVIEW_W, PREVIEW_H),
                                 background_color="#E5E7EB")]],
                      justification="center", background_color="#E5E7EB",
                      expand_x=True)]], expand_x=True)],
        [sg.Text("Value read (edit to correct)",
                 font=("Segoe UI", 10, "bold"))],
        [sg.Multiline(key="-REV_VALUE-", size=(70, 4), expand_x=True)],
        [sg.Column([[
            sg.Push(),
            sg.Button("Not legible", key="-REV_SKIP-", disabled=True, size=(14, 2),
                      button_color=(C["text"], "#E5E7EB")),
            sg.Button("Save correction", key="-REV_SAVE-", disabled=True,
                      size=(20, 2), font=("Segoe UI", 10, "bold"))]],
                  expand_x=True, pad=(0, (8, 20)))],
    ], expand_x=True, expand_y=True)

    return [[left, sg.VerticalSeparator(), right]]


def tab_reconciliation():
    return [
        [sg.Text("Reconciliation between documents",
                 font=("Segoe UI", 14, "bold"), pad=(0, (20, 4))),
         sg.Push(),
         sg.Text("0", key="-RECON_DISCREPANCIA-", text_color="#FFFFFF",
                 background_color=C["bad"], font=("Segoe UI", 10, "bold"),
                 pad=(8, 4)),
         sg.Text("discrepancies", text_color=C["muted"])],
        [sg.Text("   VERIFIED — the values match and the extraction is reliable",
                 text_color=C["ok"]),
         sg.Text("   DISCREPANCY — the values differ, with certainty",
                 text_color=C["bad"]),
         sg.Text("   NOT_VERIFIABLE — a document is missing or the reading isn't reliable",
                 text_color=C["warn"])],
        [sg.Text("Shipment"),
         sg.Combo(["All"], default_value="All", readonly=True,
                  key="-RECON_FILTER-", enable_events=True, size=(20, 1)),
         sg.Text("Verdict"),
         sg.Combo(["All", "VERIFIED", "DISCREPANCY", "NOT_VERIFIABLE"],
                  default_value="All", readonly=True,
                  key="-RECON_VERDICT-", enable_events=True, size=(18, 1))],
        [sg.Column([[sg.Table(values=[], key="-RECON_TABLE-",
                  headings=["Shipment", "Rule", "Description", "Verdict", "Detail"],
                  col_widths=[14, 6, 40, 14, 70],
                  auto_size_columns=False, justification="left", num_rows=16,
                  font=("Segoe UI", 10), header_font=("Segoe UI", 10, "bold"),
                  alternating_row_color="#F8FAFC",
                  selected_row_colors=(C["text"], C["sel"]),
                  # el scroll horizontal lo da el Column de afuera (ver
                  # nota abajo) -- la tabla NO crea su propio hsb. El
                  # vertical lo sigue manejando la tabla (su propio vsb):
                  # el alto del Treeview queda fijo en num_rows filas sin
                  # importar cuanto mida el Column que lo envuelve, asi que
                  # si se lo oculta las filas de mas quedan inalcanzables
                  # (probado aparte). El vsb del Column se oculta a mano
                  # despues de Finalize para no tener dos barras verticales.
                  vertical_scroll_only=True, hide_vertical_scroll=False,
                  enable_click_events=True,
                  expand_x=False, expand_y=False)]],
                  # Un Column con size= (SIN scrollable=True) solo recorta
                  # visualmente -- el contenido que sobra queda oculto y
                  # sin ninguna barra conectada para alcanzarlo (probado
                  # aparte: el hsb de la tabla quedaba en (0.0, 1.0), como si
                  # nada sobrara, aunque el Treeview media 1946px dentro de
                  # un Column de 1260px). scrollable=True sí cablea un
                  # scrollbar funcional al contenido interno.
                  scrollable=True, vertical_scroll_only=False,
                  size=(1260, 340), expand_x=True, expand_y=True,
                  pad=(0, (0, 24)))],
        [sg.Text("Click a row to see its full Description/Detail below "
                 "(the clicked cell is also copied).",
                 text_color=C["muted"]),
         sg.Push(),
         sg.Text("", key="-RECON_COPIED-", text_color=C["ok"])],
        [sg.Multiline("", key="-RECON_DETAIL-", size=(140, 5), disabled=True,
                      expand_x=True, font=("Segoe UI", 11), pad=(0, (3, 20)))],
    ]


# =============================================================================
# Worker
# =============================================================================
def run_extraction(folder, rules, out_q):
    """Corre en un hilo aparte: si corriera en el principal, la ventana
    quedaria congelada durante todo el OCR (12s por documento escaneado)."""
    try:
        def progress(done, total, name, result):
            out_q.put(("progress", (done, total, name, result)))
        results = extract.process_folder(folder, rules, on_progress=progress)
        out_q.put(("done", results))
    except Exception:
        out_q.put(("error", traceback.format_exc()))


# =============================================================================
# Main
# =============================================================================
def main():
    rules = yaml.safe_load(open(HERE / "rules.yaml", encoding="utf-8"))
    recon_rules = reconciliation.load_rules(HERE / "reconciliation.yaml")
    store = ReviewStore(CORRECTIONS)

    layout = [
        [sg.Column([[sg.Text("China VAT Refund", font=("Segoe UI", 17, "bold"),
                             background_color=C["surface"])],
                    [sg.Text("Extraction, review and reconciliation of "
                             "export documents",
                             font=("Segoe UI", 9), text_color=C["muted"],
                             background_color=C["surface"])]],
                   background_color=C["surface"], expand_x=True, pad=(20, 8))],
        [sg.HorizontalSeparator()],
        [sg.TabGroup([[sg.Tab("Process", tab_process()),
                       sg.Tab("Results", tab_results()),
                       sg.Tab("Review", tab_review()),
                       sg.Tab("Reconciliation", tab_reconciliation())]],
                     expand_x=True, expand_y=True, pad=(14, 12))],
        [sg.HorizontalSeparator(pad=(0, (10, 0)))],
        [sg.Push(), sg.Button("Exit", key="-EXIT-", size=(9, 1),
                              font=("Segoe UI", 10, "bold"), pad=(20, (16, 20)))],
    ]

    win = sg.Window("China VAT Refund", layout, size=(1320, 900),
                    resizable=True, finalize=True, margins=(0, 0))

    # el Column que envuelve -RECON_TABLE- trae su propio par vsb/hsb (ver
    # tab_reconciliation) -- nos quedamos solo con su hsb, el vertical ya
    # lo resuelve la tabla con el suyo. Sin esto quedan dos barras
    # verticales una al lado de la otra.
    win["-RECON_TABLE-"].ParentContainer.vsb.pack_forget()

    results, rows, cases, out_q, worker = [], [], [], queue.Queue(), None
    recon_shipments, recon_rows, recon_shown = [], [], []

    def log(msg):
        win["-LOG-"].print(msg)

    def status(text, color, pct=None):
        win["-ICON-"].update(text_color=color)
        win["-STATUS-"].update(text, text_color=color)
        if pct is not None:
            win["-PCT-"].update(f"{pct}%")
            win["-BAR-"].update(current_count=pct, max=100)

    def refresh_table():
        f = win["-FILTER-"].get()
        q = (win["-SEARCH-"].get() or "").strip().lower()
        shown = [r for r in rows
                 if (f == "All" or r[6] == f)
                 and (not q or q in " ".join(map(str, r)).lower())]
        colors = []
        for i, r in enumerate(shown):
            # row_colors de FreeSimpleGUI espera (fila, texto, fondo) -- OJO,
            # no (fila, fondo, texto), aunque el nombre de las variables lo
            # sugiera. Iba al reves y pintaba las filas con los colores
            # cambiados (fondo oscuro donde tendria que ser clarito).
            bg, fg = {"OK": ("#F0FDF4", "#166534"),
                      "OK_CORROBORATED": ("#ECFDF5", "#065F46"),
                      "CORRECTED": ("#EFF6FF", "#1D4ED8"),
                      "NOT_FOUND": ("#FEF2F2", "#991B1B"),
                      "PARTIAL": ("#FFF7ED", "#9A3412"),
                      "UNRESOLVED": ("#FFF7ED", "#9A3412"),
                      "ERROR": ("#FEF2F2", "#991B1B"),
                      }.get(r[6], ("#F9FAFB", "#4B5563"))
            colors.append((i, fg, bg))
        win["-TABLE-"].update(values=shown, row_colors=colors)

    def refresh_review():
        win["-REV_LIST-"].update(
            [f"{c['doc']} · p{c['page']} · {c['field']}" for c in cases])
        win["-REV_N-"].update(len(cases))

    def refresh_counters():
        pages = sum(len(r.get("pages", [])) for r in results)
        ok = sum(1 for r in rows if r[6] in ("OK", "OK_CORROBORATED", "CORRECTED"))
        unk = sum(1 for r in rows if r[6] in ("UNRESOLVED", "ERROR"))
        for k, v in (("-N_DOCS-", len(results)), ("-N_PAGES-", pages),
                     ("-N_OK-", ok), ("-N_REV-", len(cases)), ("-N_UNK-", unk)):
            win[k].update(v)

    def refresh_reconciliation():
        nonlocal recon_shown
        embarque = win["-RECON_FILTER-"].get()
        veredicto = win["-RECON_VERDICT-"].get()
        shown = [r for r in recon_rows
                 if (embarque == "All" or r[0] == embarque)
                 and (veredicto == "All" or r[3] == veredicto)]
        recon_shown = shown
        # (fila, texto, fondo) -- mismo detalle que en refresh_table()
        colors = [(i, *reversed({"VERIFIED": ("#F0FDF4", "#166534"),
                                 "DISCREPANCY": ("#FEF2F2", "#991B1B"),
                                 "NOT_VERIFIABLE": ("#FFEDD5", "#C2410C"),
                                 }[r[3]])) for i, r in enumerate(shown)]
        win["-RECON_TABLE-"].update(values=[_truncate_recon_row(r) for r in shown],
                                    row_colors=colors)
        # sin esto el Column con scrollable=True no se entera de que el
        # contenido cambió y el hsb queda con el rango de la carga anterior
        win["-RECON_TABLE-"].ParentContainer.contents_changed()
        n_disc = sum(1 for r in recon_rows if r[3] == "DISCREPANCY")
        win["-RECON_DISCREPANCIA-"].update(n_disc)

    def show_case(c):
        win["-REV_EMPTY-"].update(visible=False)
        win["-REV_DOC-"].update(c["doc"])
        win["-REV_PAGE-"].update(c["page"])
        win["-REV_CONF-"].update(f"{c['conf']:.2f}")
        win["-REV_FIELD-"].update(c["field"])
        win["-REV_ZH-"].update(f"({c['desc_zh']})" if c["desc_zh"] else "")
        win["-REV_STATUS-"].update(c["status"])
        win["-REV_VALUE-"].update(fmt_value(c["value"]))
        img = render_region(c["abs_path"], c["page"], c["origin"],
                            c["render_scale"], c.get("rotation", 0))
        win["-REV_IMG-"].update(data=img)
        win["-REV_SAVE-"].update(disabled=False)
        win["-REV_SKIP-"].update(disabled=False)

    while True:
        ev, val = win.read(timeout=120)
        if ev in (sg.WIN_CLOSED, "-EXIT-"):
            break

        # --- cola del worker
        while not out_q.empty():
            kind, payload = out_q.get()
            if kind == "progress":
                done, total, name, res = payload
                pct = int(done / max(total, 1) * 100)
                if res is None:
                    status(f"Processing {name}…", C["primary"], pct)
                else:
                    mark = "OK " if res["ok"] else "ERR"
                    log(f"[{done}/{total}] {mark} {name}  ({res['secs']}s)")
                    if not res["ok"]:
                        # sin esto el log dice 'ERR' y nada mas, que no
                        # alcanza para saber que pasó
                        log(f"        {res.get('error', 'no detail')}")
                        if res.get("traceback"):
                            log(res["traceback"])
                    status(f"{done}/{total} processed", C["primary"], pct)
            elif kind == "done":
                results = payload
                n = store.apply(results)
                if n:
                    log(f"Reapplied {n} saved corrections.")
                rows = flatten(results)
                cases = build_review_cases(results)
                recon_shipments = reconciliation.run(results, recon_rules,
                                                     rules["doc_types"])
                recon_rows = flatten_reconciliation(recon_shipments)
                refresh_table()
                refresh_review()
                refresh_counters()
                win["-RECON_FILTER-"].update(
                    values=["All"] + sorted({r[0] for r in recon_rows}))
                refresh_reconciliation()
                win["-PROCESS-"].update(disabled=False, text="Process")
                win["-EXPORT-"].update(disabled=not rows)
                n_disc = sum(1 for r in recon_rows if r[3] == "DISCREPANCY")
                status(f"Done — {len(results)} documents, {len(cases)} to review, "
                       f"{n_disc} discrepancies", C["ok"], 100)
                log("Done.")
            elif kind == "error":
                log(payload)
                win["-PROCESS-"].update(disabled=False, text="Process")
                status("Error during processing", C["bad"], 0)

        # --- eventos
        if ev == "-PROCESS-":
            folder = Path(val["-FOLDER-"])
            if not folder.is_dir():
                sg.popup_error("The folder doesn't exist:", str(folder))
                continue
            if not any(p.suffix.lower() in extract.PROCESSORS
                      for p in folder.rglob("*")):
                sg.popup_error("No PDFs or xlsx files in:", str(folder))
                continue
            win["-LOG-"].update("")
            win["-PROCESS-"].update(disabled=True, text="Processing…")
            status("Starting…", C["primary"], 0)
            worker = threading.Thread(target=run_extraction,
                                      args=(folder, rules, out_q), daemon=True)
            worker.start()

        elif ev in ("-FILTER-", "-SEARCH-"):
            refresh_table()

        elif ev in ("-RECON_FILTER-", "-RECON_VERDICT-"):
            refresh_reconciliation()

        elif isinstance(ev, tuple) and ev[0] == "-RECON_TABLE-":
            _, _, (row, col) = ev
            if row is not None and row >= 0 and col is not None and col >= 0 \
                    and row < len(recon_shown):
                value = str(recon_shown[row][col])
                sg.clipboard_set(value)
                shown_value = value if len(value) <= 60 else value[:57] + "…"
                win["-RECON_COPIED-"].update(f"Copied: {shown_value}")
                # el panel de abajo muestra la fila entera, no solo la celda
                # clickeada -- asi no hace falta acertarle a la columna exacta
                # (Description o Detail) para leer el texto completo
                full_row = recon_shown[row]
                win["-RECON_DETAIL-"].update(
                    f"Description: {full_row[2]}\n\nDetail: {full_row[4]}")

        elif ev == "-REV_LIST-" and val["-REV_LIST-"]:
            idx = win["-REV_LIST-"].get_indexes()
            if idx:
                show_case(cases[idx[0]])

        elif ev in ("-REV_SAVE-", "-REV_SKIP-"):
            idx = win["-REV_LIST-"].get_indexes()
            if not idx:
                sg.popup_error("Select a case first.")
                continue
            c = cases[idx[0]]
            if ev == "-REV_SAVE-":
                new = val["-REV_VALUE-"].strip()
                if not new:
                    sg.popup_error("The value can't be empty. "
                                   "Use 'Not legible' if the document can't be read.")
                    continue
                store.set(c["doc"], c["page"], c["field"], new,
                          fmt_value(c["value"]))
                log(f"Corrected {c['field']} in {c['doc']} p{c['page']}: {new}")
            else:
                store.set(c["doc"], c["page"], c["field"], None,
                          fmt_value(c["value"]), reason="ilegible")
                log(f"Marked not legible: {c['field']} in {c['doc']} p{c['page']}")

            # reflejar en la tabla sin reprocesar
            for r in rows:
                if r[0] == c["doc"] and r[1] == c["page"] and r[3] == c["field"]:
                    r[4] = val["-REV_VALUE-"].strip() if ev == "-REV_SAVE-" else ""
                    r[6] = "CORRECTED" if ev == "-REV_SAVE-" else "ILLEGIBLE"
                    break
            cases.pop(idx[0])
            refresh_table()
            refresh_review()
            refresh_counters()
            if cases:
                show_case(cases[min(idx[0], len(cases) - 1)])
            else:
                for k in ("-REV_DOC-", "-REV_PAGE-", "-REV_CONF-", "-REV_FIELD-",
                          "-REV_ZH-", "-REV_STATUS-", "-REV_VALUE-"):
                    win[k].update("")
                win["-REV_IMG-"].update(data=None)
                win["-REV_SAVE-"].update(disabled=True)
                win["-REV_SKIP-"].update(disabled=True)
                win["-REV_EMPTY-"].update("No pending cases left.",
                                          text_color=C["ok"], visible=True)

        elif ev == "-EXPORT-":
            path = sg.popup_get_file("Save as", save_as=True,
                                     default_extension=".xlsx",
                                     file_types=(("Excel", "*.xlsx"),))
            if path:
                try:
                    export_excel(rows, path)
                    sg.popup("Exported:", path)
                except Exception as e:
                    sg.popup_error("Could not export:", str(e))

    win.close()


def export_excel(rows, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    if not str(path).lower().endswith(".xlsx"):
        path = f"{path}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Extraction"
    head = ["File", "Page", "Document type", "Field", "Value",
            "Confidence", "Status"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    for col, w in zip("ABCDEFG", (34, 8, 26, 24, 46, 11, 18)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(path)


if __name__ == "__main__":
    main()