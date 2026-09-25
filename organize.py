"""
Salida separada por embarque: una carpeta por exportacion con copia de sus
documentos y un reporte Excel (resumen, checklist, reconciliacion y datos
extraidos).

    output/exports/<fecha_hora>/<SO>/documents/...
    output/exports/<fecha_hora>/<SO>/<SO>_report.xlsx
    output/exports/<fecha_hora>/_unassigned/...      (si hay paginas sin asignar)

Cada corrida va a su propia carpeta con fecha y hora: nunca se pisa ni se
borra una salida anterior. Los originales no se tocan (se copian).

Los logs compartidos (per_row en rules.yaml: Shipping Order, Tax Filing
Directory) NO se copian a cada carpeta -- traen filas de todos los
embarques. Sus filas de ESTE embarque aparecen en la hoja "Extracted data".

Correr solo:  python organize.py [carpeta]      (por defecto samples/)
"""

import re
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HERE = Path(__file__).resolve().parent
UNASSIGNED_DIR = "_unassigned"

# (fondo, texto) -- mismos colores que la GUI
STATUS_COLORS = {
    "PRESENT": ("F0FDF4", "166534"), "VERIFIED": ("F0FDF4", "166534"),
    "OK": ("F0FDF4", "166534"), "OK_CORROBORATED": ("F0FDF4", "166534"),
    "CORRECTED": ("EFF6FF", "1D4ED8"),
    "MISSING": ("FEF2F2", "991B1B"), "DISCREPANCY": ("FEF2F2", "991B1B"),
    "NOT_FOUND": ("FEF2F2", "991B1B"), "CONFLICT": ("FEF2F2", "991B1B"),
    "MANUAL": ("FFEDD5", "C2410C"), "NOT_VERIFIABLE": ("FFEDD5", "C2410C"),
    "PARTIAL": ("FFEDD5", "C2410C"),
}


def _safe_name(text, limit=60):
    """Nombre de carpeta valido en Windows a partir de la etiqueta del embarque."""
    name = re.sub(r'[<>:"/\\|?*\s]+', "_", str(text)).strip("._")
    return (name or "shipment")[:limit]


def _fmt(v):
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, list):
        return ", ".join(map(str, v))
    if isinstance(v, dict):
        return ", ".join(f"{k}={x}" for k, x in v.items())
    return "" if v is None else str(v)


def _sheet(wb, title, head, rows, widths, status_col=None, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(ord("A") + i)].width = w
    for row in ws.iter_rows(min_row=2):
        colors = STATUS_COLORS.get(str(row[status_col].value)) if status_col is not None else None
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if colors:
                cell.fill = PatternFill("solid", fgColor=colors[0])
                cell.font = Font(color=colors[1])
    return ws


def _sources(sources):
    by_file = {}
    for s in sources:
        by_file.setdefault(s["file"], []).append(s["page"])
    return " | ".join(f"{f} ({', '.join(map(str, sorted(p)))})"
                      for f, p in by_file.items())


def _leg(leg):
    ref = leg.get("ref") or f"item {leg.get('item')} (no extractor)"
    if leg["status"] == "MISSING":
        return f"{ref}: not available"
    return f"{ref}: {_fmt(leg['value'])} [{leg['status']}]"


def _field_rows(records, pages_by_key):
    rows = []
    for rec in sorted(records, key=lambda r: (r["file"], r["page"])):
        page = pages_by_key.get((rec["file"], rec["page"]))
        if not page:
            continue
        if not page["fields"]:
            rows.append([rec["file"], rec["page"], rec["doc_type"], "", "", "", "OK"])
        for fname, f in page["fields"].items():
            rows.append([rec["file"], rec["page"], rec["doc_type"], fname,
                         _fmt(f["value"]), round(f["conf"], 2), f["status"]])
    return rows


def write_report(path, shipment, source_folder, pages_by_key, log_files):
    ck = shipment["checklist"]
    count = lambda items, key, val: sum(1 for i in items if i[key] == val)
    wb = Workbook()
    summary = [
        ["Shipment", shipment["shipment"]],
        ["Generated", datetime.now().strftime("%Y-%m-%d %H:%M")],
        ["Source folder", str(source_folder)],
        ["Documents present", count(ck, "status", "PRESENT")],
        ["Documents missing", count(ck, "status", "MISSING")],
        ["Documents to check by hand", count(ck, "status", "MANUAL")],
        ["Rules verified", count(shipment["rules"], "verdict", "VERIFIED")],
        ["Rules with discrepancy", count(shipment["rules"], "verdict", "DISCREPANCY")],
        ["Rules not verifiable", count(shipment["rules"], "verdict", "NOT_VERIFIABLE")],
        ["Warnings", "\n".join(shipment["warnings"]) or "none"],
        ["Shared logs (not copied)",
         "\n".join(log_files) + "\nTheir rows for this shipment are in 'Extracted data'."
         if log_files else "none"],
    ]
    _sheet(wb, "Summary", ["Item", "Value"], summary, (30, 90), first=True)
    _sheet(wb, "Checklist", ["Code", "Document", "Status", "Found in"],
           [[i["code"], i["name"], i["status"], _sources(i["sources"])] for i in ck],
           (10, 60, 12, 70), status_col=2)
    _sheet(wb, "Reconciliation", ["Rule", "Description", "Verdict", "Detail"],
           [[r["id"], r["desc"], r["verdict"], "\n".join(_leg(l) for l in r["legs"])]
            for r in shipment["rules"]],
           (6, 60, 16, 80), status_col=2)
    _sheet(wb, "Extracted data",
           ["File", "Page / row", "Document type", "Field", "Value", "Confidence", "Status"],
           _field_rows(shipment["records"], pages_by_key),
           (34, 10, 24, 24, 46, 11, 18), status_col=6)
    wb.save(path)


def export_by_shipment(results, recon, doc_types, source_folder, dest_root=None):
    """
    Arma la salida por embarque. Devuelve la carpeta de la corrida.
    results = extract.process_folder(), recon = reconciliation.run().
    """
    dest_root = Path(dest_root) if dest_root else HERE / "output" / "exports"
    run_dir = dest_root / datetime.now().strftime("%Y-%m-%d_%H%M%S")
    run_dir.mkdir(parents=True, exist_ok=False)

    abs_of = {r["path"]: r.get("abs_path") for r in results}
    pages_by_key = {(r["path"], p["page"]): p
                    for r in results if r.get("ok") for p in r["pages"]}
    is_log = lambda dt: bool(doc_types.get(dt, {}).get("per_row"))

    def copy(files, docs_dir):
        docs_dir.mkdir(parents=True, exist_ok=True)
        for f in files:
            src = abs_of.get(f)
            if src and Path(src).is_file():
                # la estructura de subcarpetas de origen se aplana; si dos
                # archivos se llaman igual, el segundo lleva el prefijo de su
                # carpeta para no pisar al primero
                dst = docs_dir / Path(f).name
                if dst.exists():
                    dst = docs_dir / _safe_name(str(Path(f).parent) + "_" + Path(f).name, 150)
                shutil.copy2(src, dst)

    for s in recon["shipments"]:
        d = run_dir / _safe_name(s["shipment"])
        own = sorted({r["file"] for r in s["records"] if not is_log(r["doc_type"])})
        logs = sorted({r["file"] for r in s["records"] if is_log(r["doc_type"])})
        copy(own, d / "documents")
        write_report(d / f"{_safe_name(s['shipment'])}_report.xlsx", s,
                     source_folder, pages_by_key, logs)

    if recon["unassigned"]:
        d = run_dir / UNASSIGNED_DIR
        copy(sorted({u["file"] for u in recon["unassigned"]}), d / "documents")
        wb = Workbook()
        _sheet(wb, "Unassigned pages", ["File", "Page / row", "Document type"],
               [[u["file"], u["page"], u["doc_type"]] for u in recon["unassigned"]],
               (50, 10, 30), first=True)
        wb.save(d / "unassigned_pages.xlsx")
    return run_dir


if __name__ == "__main__":
    import sys

    import yaml

    import extract
    import reconciliation

    folder = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "samples"
    rules_yaml = yaml.safe_load(open(HERE / "rules.yaml", encoding="utf-8"))
    results = extract.process_folder(folder, rules_yaml)
    recon = reconciliation.run(results, doc_types=rules_yaml["doc_types"])
    out = export_by_shipment(results, recon, rules_yaml["doc_types"], folder)
    print(f"\nSalida por embarque: {out}")
    for p in sorted(out.iterdir()):
        print(f"  {p.name}/  ({len(list((p / 'documents').glob('*')))} documentos)")
